# -*- coding: utf-8 -*-
"""
=====================================================================
 CASO 1 - RPA sobre METEOSTREAM Clima Chile
 Sitio: https://mantistcy.cl/clima/
=====================================================================

 Que hace este robot (4 acciones distintas sobre el sitio):

   1) NAVEGAR  -> abre la portada de MeteoStream y espera a que la
                  telemetria dinamica termine de cargar.
   2) EXTRAER  -> hace scraping de las tarjetas de estaciones
                  (temperatura, humedad, MP 2.5, condicion) y guarda
                  todo en archivos CSV y XLSX + capturas de pantalla.
                  Compara MP 2.5 y lluvia con el maestro de umbrales.
                  Tambien descarga el archivo original de telemetria TXT.
   3) CONSULTAR -> hace clic en Consultar Sismos y captura el resultado.
   4) CARGAR   -> hace clic en avisos.php y llena el formulario de alertas
                  con la estacion mas contaminada que encontro.
=====================================================================
"""

import csv
import math
import os
import re
import sys
import time
from datetime import datetime
from tempfile import mkdtemp
from pathlib import Path
import subprocess

# ---------------------------------------------------------------
# AUTO-INSTALACION DE DEPENDENCIAS FALTANTES
# ---------------------------------------------------------------
def instalar_requerimientos():
    try:
        import selenium
        import openpyxl
    except ImportError:
        print("=" * 60)
        print(" [!] Componentes faltantes detectados.")
        print(" [!] Iniciando instalacion automatica de dependencias...")
        print("=" * 60)
        ruta_req = os.path.join(os.path.dirname(os.path.abspath(__file__)), "requirements.txt")
        if os.path.exists(ruta_req):
            subprocess.check_call([sys.executable, "-m", "pip", "install", "-r", ruta_req])
        else:
            subprocess.check_call([sys.executable, "-m", "pip", "install", "selenium>=4.25", "openpyxl>=3.1"])
        print(" [OK] Dependencias instaladas correctamente.\n")

instalar_requerimientos()

from openpyxl import Workbook, load_workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

try:
    sys.stdout.reconfigure(encoding="utf-8")
except Exception:
    pass

from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.support.ui import WebDriverWait, Select
from selenium.webdriver.support import expected_conditions as EC
from selenium.common.exceptions import (
    TimeoutException,
    NoSuchElementException,
    WebDriverException,
)

# ---------------------------------------------------------------
# 1. CONFIGURACION
# ---------------------------------------------------------------

# Navegador a usar: "chrome", "firefox" o "edge"
NAVEGADOR = "firefox"

# URLs del sitio objetivo
URL_PORTADA = "https://mantistcy.cl/clima/"

# True  = llena el formulario Y lo envia (publica el aviso)
# False = solo llena el formulario y saca la captura, sin enviar
PUBLICAR_AVISO = True

# --- TIEMPOS DE ESPERA (en segundos) ---------------------------
TIEMPO_MAXIMO = 60      # tope de espera por un elemento (WebDriverWait)
PAUSA_CARGA = 6         # pausa extra despues de cargar la portada
PAUSA_PASO = 3          # pausa entre un paso y otro
MIN_ESTACIONES = 10     # estaciones que deben existir antes de leer

# Umbrales de MP 2.5 en microgramos/m3 (referencia norma chilena diaria)
UMBRAL_CRITICO = 50
UMBRAL_MODERADO = 25

# --- DONDE SE GUARDAN LAS EVIDENCIAS ---------------------------
# Por defecto se crea la carpeta "evidencias" al lado de este archivo.

CARPETA_BASE = os.path.dirname(os.path.abspath(__file__))
ARCHIVO_UMBRALES = os.path.join(CARPETA_BASE, "maestro_umbrales_clima.xlsx")

CARPETA = os.path.join(CARPETA_BASE, "evidencias")
os.makedirs(CARPETA, exist_ok=True)


# ---------------------------------------------------------------
# 2. FUNCIONES AUXILIARES
# ---------------------------------------------------------------

def esperar(segundos, motivo):
    """Pausa controlada que ademas avisa en pantalla por que espera."""
    print("   ... esperando " + str(segundos) + "s (" + motivo + ")")
    time.sleep(segundos)


def crear_driver(navegador):
    """Abre el navegador buscando el preferido y con alternativas (fallback)."""
    navegador = navegador.lower().strip()
    descargas = mkdtemp(prefix="descarga_", dir=CARPETA)
    
    # Lista de navegadores a intentar, priorizando el elegido
    opciones_nav = ["chrome", "edge", "firefox"]
    if navegador in opciones_nav:
        opciones_nav.remove(navegador)
        orden = [navegador] + opciones_nav
    else:
        orden = opciones_nav
        print(f"   [!] Navegador '{navegador}' desconocido. Intentando auto-deteccion.")

    driver = None
    nav_exitoso = None
    
    for nav in orden:
        try:
            print(f"   Intentando iniciar navegador: {nav}...")
            if nav == "chrome":
                opciones = webdriver.ChromeOptions()
                opciones.add_argument("--window-size=1400,900")
                opciones.add_experimental_option("prefs", {
                    "download.default_directory": descargas,
                    "download.prompt_for_download": False,
                    "download.directory_upgrade": True,
                })
                driver = webdriver.Chrome(options=opciones)
                
            elif nav == "firefox":
                opciones = webdriver.FirefoxOptions()
                opciones.add_argument("--width=1400")
                opciones.add_argument("--height=900")
                opciones.set_preference("browser.download.folderList", 2)
                opciones.set_preference("browser.download.dir", descargas)
                opciones.set_preference("browser.download.useDownloadDir", True)
                opciones.set_preference("browser.helperApps.neverAsk.saveToDisk", "text/plain,application/octet-stream")
                driver = webdriver.Firefox(options=opciones)
                
            elif nav == "edge":
                opciones = webdriver.EdgeOptions()
                opciones.add_argument("--window-size=1400,900")
                opciones.add_experimental_option("prefs", {
                    "download.default_directory": descargas,
                    "download.prompt_for_download": False,
                    "download.directory_upgrade": True,
                })
                driver = webdriver.Edge(options=opciones)
            
            # Si llego aca es porque pudo iniciarlo con exito
            nav_exitoso = nav
            break
            
        except Exception as e:
            print(f"   [!] No se encontro o fallo {nav}. Detalle: {e}")
            continue

    if driver is None:
        import shutil
        shutil.rmtree(descargas, ignore_errors=True)
        raise RuntimeError("No se encontro NINGUN navegador compatible (Chrome, Edge, Firefox) en esta maquina.")

    # Tiempo maximo que el navegador espera a que la pagina cargue entera
    print(f"   [OK] Navegador cargado exitosamente: {nav_exitoso}")
    driver.set_page_load_timeout(TIEMPO_MAXIMO)
    driver.telemetria_descargas = descargas
    driver.nombre_navegador_usado = nav_exitoso
    return driver


def capturar(driver, nombre):
    """Guarda una captura de pantalla como evidencia del proceso."""
    marca = datetime.now().strftime("%Y%m%d_%H%M%S")
    ruta = os.path.join(CARPETA, marca + "_" + nombre + ".png")
    driver.save_screenshot(ruta)
    print("   [evidencia] " + ruta)
    return ruta


def a_numero(texto):
    """
    Validacion de datos: convierte "16.8 C" o "10 ug/m3" en un numero.
    Si el texto no trae ningun numero devuelve None (dato invalido).
    """
    encontrado = re.search(r"-?\d+[.,]?\d*", texto)
    if encontrado is None:
        return None
    try:
        return float(encontrado.group().replace(",", "."))
    except ValueError:
        return None


def clasificar_aire(mp25):
    """Clasifica la calidad del aire. Ejemplo claro de if / elif / else."""
    if mp25 is None:
        return "SIN DATO"
    elif mp25 >= UMBRAL_CRITICO:
        return "CRITICO"
    elif mp25 >= UMBRAL_MODERADO:
        return "MODERADO"
    else:
        return "BUENO"


def contar_estaciones(driver):
    """Cuenta cuantas tarjetas de estacion hay ya dibujadas en la pagina."""
    tarjetas = driver.find_elements(By.CSS_SELECTOR, "#tiempo-real div.bd-card.h-100")
    total = 0
    for t in tarjetas:                       # ciclo for
        # Solo cuentan las tarjetas que tienen titulo: esas son estaciones
        if len(t.find_elements(By.TAG_NAME, "h5")) > 0:
            total = total + 1
    return total


def esperar_datos_completos(driver):
    """
    Espera a que la telemetria termine de cargar, en dos etapas:

      Etapa 1: espera hasta que existan al menos MIN_ESTACIONES tarjetas.
      Etapa 2: CICLO WHILE que vuelve a contar cada 2 segundos hasta que
               el numero deje de cambiar. Asi el robot no lee la pagina
               a medio dibujar.
    """
    # Etapa 1: espera activa con tope de TIEMPO_MAXIMO segundos
    WebDriverWait(driver, TIEMPO_MAXIMO).until(
        lambda d: contar_estaciones(d) >= MIN_ESTACIONES
    )

    # Etapa 2: se espera a que la cantidad se estabilice
    anterior = -1
    actual = contar_estaciones(driver)
    intentos = 0

    while anterior != actual and intentos < 10:   # CICLO WHILE
        anterior = actual
        time.sleep(2)
        actual = contar_estaciones(driver)
        intentos = intentos + 1
        print("   ... telemetria cargando: " + str(actual) + " estaciones")

    print("   Telemetria estable en " + str(actual) + " estaciones.")
    return actual


# ---------------------------------------------------------------
# 3. ACCIONES 1 y 2: NAVEGAR + EXTRAER LA TELEMETRIA
# ---------------------------------------------------------------

def extraer_estaciones(driver):
    """Entra a la portada y devuelve una lista de diccionarios,
    uno por cada estacion meteorologica publicada."""

    print("\n[1] Navegando a la portada de MeteoStream...")
    driver.get(URL_PORTADA)

    # Espera INTELIGENTE: en vez de un time.sleep a ciegas, el robot
    # espera a que la telemetria dinamica termine de cargarse.
    esperar_datos_completos(driver)

    # Pausa adicional para que terminen de dibujarse graficos e iconos
    esperar(PAUSA_CARGA, "que terminen de dibujarse los graficos")

    print("   Portada cargada correctamente.")
    capturar(driver, "01_portada")

    print("\n[2] Extrayendo telemetria de las estaciones...")
    tarjetas = driver.find_elements(By.CSS_SELECTOR, "#tiempo-real div.bd-card.h-100")
    print("   Se encontraron " + str(len(tarjetas)) + " bloques en el panel.")

    estaciones = []

    # CICLO DE REPETICION: recorre una por una todas las tarjetas del panel
    for tarjeta in tarjetas:

        # CONDICIONAL 1: no todas las tarjetas son estaciones (algunas son
        # graficos). Si no tiene titulo <h5>, se salta y sigue con la otra.
        titulos = tarjeta.find_elements(By.TAG_NAME, "h5")
        if len(titulos) == 0:
            continue

        datos = {
            "estacion": titulos[0].text.strip(),
            "zona": "",
            "temperatura_c": None,
            "humedad_pct": None,
            "mp25": None,
            "lluvia_mm": None,
            "condicion": "",
        }

        # La zona (Norte / Centro / Sur) viene en el primer <span> de la tarjeta
        try:
            datos["zona"] = tarjeta.find_element(By.TAG_NAME, "span").text.strip()
        except NoSuchElementException:
            datos["zona"] = "SIN ZONA"

        # Cada tarjeta tiene 4 recuadros: Temp, Humedad, MP 2.5 y Condicion.
        # Se leen por su etiqueta y no por posicion, para que el robot no se
        # rompa si el sitio cambia el orden de los recuadros.
        recuadros = tarjeta.find_elements(By.CSS_SELECTOR, "div.row.g-2 div.col-6")

        for recuadro in recuadros:  # ciclo anidado sobre los 4 recuadros
            etiqueta = recuadro.find_element(By.TAG_NAME, "small").text.strip().upper()
            valor = recuadro.find_element(By.TAG_NAME, "span").text.strip()

            # CONDICIONAL 2: segun la etiqueta, el valor se guarda donde toca
            if "TEMP" in etiqueta:
                datos["temperatura_c"] = a_numero(valor)
            elif "HUMEDAD" in etiqueta:
                datos["humedad_pct"] = a_numero(valor)
            elif "MP" in etiqueta:
                datos["mp25"] = a_numero(valor)
            elif "CONDICI" in etiqueta:
                datos["condicion"] = valor
            else:
                pass  # etiqueta desconocida: se ignora

        # La lluvia aparece en el pie de la tarjeta, fuera de los 4 recuadros.
        lluvia = tarjeta.find_elements(
            By.XPATH, './/span[starts-with(normalize-space(.), "Lluvia:")]'
        )
        if lluvia:
            datos["lluvia_mm"] = a_numero(lluvia[0].text)

        # Se agrega la clasificacion calculada por el robot
        datos["calidad_aire"] = clasificar_aire(datos["mp25"])
        estaciones.append(datos)

        print(
            "   - " + datos["estacion"].ljust(20)
            + " " + str(datos["temperatura_c"]) + " C"
            + " | HR " + str(datos["humedad_pct"]) + "%"
            + " | MP2.5 " + str(datos["mp25"])
            + " -> " + datos["calidad_aire"]
        )

    return estaciones


def descargar_telemetria(driver):
    """Hace clic en Telemetria Raw y espera a que el navegador guarde el TXT."""
    enlace = WebDriverWait(driver, TIEMPO_MAXIMO).until(
        EC.element_to_be_clickable(
            (By.PARTIAL_LINK_TEXT, "Raw (.TXT)")
        )
    )
    descargas = driver.telemetria_descargas
    anteriores = set(os.listdir(descargas))
    print("\n   Descargando telemetria raw TXT con un clic...")
    enlace.click()

    def descarga_completa(_):
        nombres = set(os.listdir(descargas)) - anteriores
        # Chrome/Edge usan .crdownload; Firefox usa .part durante la descarga.
        if any(n.endswith((".crdownload", ".part", ".tmp")) for n in nombres):
            return False
        for nombre in nombres:
            ruta_descarga = os.path.join(descargas, nombre)
            if nombre.lower().endswith(".txt") and os.path.isfile(ruta_descarga):
                if os.path.getsize(ruta_descarga) > 0:
                    return ruta_descarga
        return False

    origen = WebDriverWait(driver, TIEMPO_MAXIMO, poll_frequency=0.5).until(
        descarga_completa, message="No se completo la descarga de telemetria TXT."
    )
    marca = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
    ruta = os.path.join(CARPETA, marca + "_stream_clima_chile.txt")
    os.rename(origen, ruta)
    print("   [telemetria TXT] " + ruta + " (" + str(os.path.getsize(ruta)) + " bytes)")
    return ruta


def guardar_csv(estaciones):
    """Integracion con planilla local: deja los datos en un CSV que se
    abre directo en Excel."""
    ruta = os.path.join(CARPETA, "telemetria_meteostream.csv")
    columnas = [
        "estacion", "zona", "temperatura_c", "humedad_pct",
        "mp25", "lluvia_mm", "condicion", "calidad_aire", "estado_alerta",
    ]

    with open(ruta, "w", newline="", encoding="utf-8-sig") as archivo:
        escritor = csv.DictWriter(archivo, fieldnames=columnas, delimiter=";")
        escritor.writeheader()
        for fila in estaciones:          # ciclo para escribir cada estacion
            escritor.writerow(fila)

    print("\n   [planilla] " + ruta + "  (" + str(len(estaciones)) + " filas)")
    return ruta


# ---------------------------------------------------------------
# EXPORTACION EXCEL Y COMPARACION CON EL MAESTRO
# ---------------------------------------------------------------

def cargar_umbrales(ruta):
    """Lee parametro/umbral en la hoja Umbrales; nunca inventa valores faltantes."""
    if not Path(ruta).is_file():
        raise FileNotFoundError(
            f"Falta {ruta}. Use la hoja Umbrales con columnas parametro y umbral, "
            "y filas mp25 y lluvia_mm."
        )
    libro = load_workbook(ruta, read_only=True, data_only=True)
    try:
        if "Umbrales" not in libro.sheetnames:
            raise ValueError("El maestro debe contener la hoja Umbrales.")
        filas = libro["Umbrales"].iter_rows(values_only=True)
        encabezados = [str(v).strip().lower() for v in next(filas, ())]
        if any(encabezados.count(c) != 1 for c in ("parametro", "umbral")):
            raise ValueError("El maestro requiere columnas unicas parametro y umbral.")
        parametro_col = encabezados.index("parametro")
        umbral_col = encabezados.index("umbral")
        umbrales = {}
        for fila in filas:
            parametro = str(fila[parametro_col]).strip().lower()
            if parametro not in ("mp25", "lluvia_mm"):
                continue
            valor = fila[umbral_col]
            if parametro in umbrales:
                raise ValueError(f"Parametro duplicado en el maestro: {parametro}")
            if (isinstance(valor, bool) or not isinstance(valor, (int, float))
                    or not math.isfinite(valor) or valor < 0):
                raise ValueError(f"Umbral invalido para {parametro}: use un numero >= 0.")
            umbrales[parametro] = float(valor)
        if set(umbrales) != {"mp25", "lluvia_mm"}:
            raise ValueError("El maestro debe definir mp25 y lluvia_mm.")
        return umbrales
    finally:
        libro.close()


def evaluar_alerta(estacion, umbrales):
    """OR: basta superar un umbral; faltar un dato no equivale a cero."""
    mp25 = estacion.get("mp25")
    lluvia = estacion.get("lluvia_mm")
    if ((mp25 is not None and mp25 > umbrales["mp25"])
            or (lluvia is not None and lluvia > umbrales["lluvia_mm"])):
        return "ALERTA"
    if mp25 is None or lluvia is None:
        return "SIN DATO"
    return "SIN ALERTA"


def guardar_excel(estaciones, umbrales, carpeta, fuente):
    """Exporta datos, umbrales usados y una formula de alerta editable en Excel."""
    fecha = datetime.now()
    ruta = Path(carpeta) / f"clima_consolidado_{fecha:%Y%m%d_%H%M%S_%f}.xlsx"
    libro = Workbook()
    hoja = libro.active
    hoja.title = "Consolidado"
    hoja.append(["Clima consolidado"])
    hoja.append(["Fecha de extraccion", fecha])
    hoja["B2"].number_format = "yyyy-mm-dd hh:mm:ss"
    hoja.append(["Fuente", fuente])
    hoja.append(["Umbral MP 2.5 (µg/m³)", umbrales["mp25"],
                 "Umbral lluvia (mm)", umbrales["lluvia_mm"]])
    hoja.append(["Regla: MP 2.5 > umbral OR lluvia > umbral. Maestro: maestro_umbrales_clima.xlsx"])
    hoja.append([])
    hoja.append(["Estacion", "Zona", "Temperatura (°C)", "Humedad (%)",
                 "MP 2.5 (µg/m³)", "Lluvia (mm)", "Condicion", "Calidad aire", "Estado alerta"])
    claves = ("estacion", "zona", "temperatura_c", "humedad_pct", "mp25",
              "lluvia_mm", "condicion", "calidad_aire")
    for numero, estacion in enumerate(estaciones, start=8):
        hoja.append([estacion.get(clave) for clave in claves])
        # Texto del sitio se conserva como texto, incluso si comienza con '='.
        for celda in hoja[numero]:
            if isinstance(celda.value, str):
                celda.data_type = "s"
        hoja.cell(numero, 9,
            f'=IF(OR(AND(ISNUMBER(E{numero}),E{numero}>$B$4),'
            f'AND(ISNUMBER(F{numero}),F{numero}>$D$4)),"ALERTA",'
            f'IF(COUNT(E{numero}:F{numero})<2,"SIN DATO","SIN ALERTA"))')
        for columna in range(3, 7):
            hoja.cell(numero, columna).number_format = "0.0"

    for fila in hoja:
        for celda in fila:
            celda.font = Font(name="Arial", size=11)
    hoja["A1"].font = Font(name="Arial", size=14, bold=True)
    for celda in hoja[7]:
        celda.font = Font(name="Arial", size=11, bold=True, color="FFFFFF")
        celda.fill = PatternFill("solid", fgColor="334155")
        celda.alignment = Alignment(wrap_text=True, vertical="center")
    hoja.row_dimensions[7].height = 32
    for i, ancho in enumerate((29, 23, 24, 20, 21, 19, 30, 18, 19), start=1):
        hoja.column_dimensions[get_column_letter(i)].width = ancho
    hoja.freeze_panes = "C8"
    hoja.sheet_view.showGridLines = False
    hoja.auto_filter.ref = f"A7:I{max(7, hoja.max_row)}"
    if estaciones:
        hoja.conditional_formatting.add(
            f"I8:I{hoja.max_row}", FormulaRule(
                formula=['I8="ALERTA"'],
                fill=PatternFill("solid", fgColor="FEE2E2"),
                font=Font(color="991B1B", bold=True),
            )
        )
        
    # --- MARCA DE AGUA ---
    ultima_fila = hoja.max_row + 2
    hoja.cell(ultima_fila, 1, "Hecho por Fabiola villagra Ignacio gonzales y Francisca Sanhueza").font = Font(
        name="Arial", size=10, italic=True, color="888888"
    )
    
    libro.save(ruta)
    libro.close()
    print(f"   [Excel consolidado] {ruta} ({len(estaciones)} estaciones)")
    return str(ruta)


# ---------------------------------------------------------------
# 4. ACCIONES 3 y 4: CONSULTAR SISMOS Y CARGAR AVISOS
# ---------------------------------------------------------------

def consultar_sismos(driver):
    """Abre la consulta mediante un clic y vuelve para continuar con los avisos."""
    print("\n[3] Consultando sismos mediante un clic...")
    # Se usa href porque el sitio repite el id sismos-tab en otro enlace.
    WebDriverWait(driver, TIEMPO_MAXIMO).until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, 'a[href="obtener_sismos.php"]'))
    ).click()
    WebDriverWait(driver, TIEMPO_MAXIMO).until(
        EC.url_contains("/obtener_sismos.php")
    )
    WebDriverWait(driver, TIEMPO_MAXIMO).until(
        lambda d: d.find_element(By.TAG_NAME, "body").text.strip()
    )
    capturar(driver, "02_sismos")
    esperar(PAUSA_PASO, "ver el resultado de la consulta de sismos")

    # La respuesta es JSON y no tiene menu: volver con el historial del navegador.
    driver.back()
    WebDriverWait(driver, TIEMPO_MAXIMO).until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, 'a[href="avisos.php"]'))
    )


# El formulario pide Region, pero el panel entrega ciudades.
# Este diccionario traduce ciudad -> region oficial.
REGIONES = {
    "Arica": "Arica y Parinacota",
    "Antofagasta": "Antofagasta",
    "Copiapo": "Atacama",
    "Copiap\u00f3": "Atacama",
    "Concon": "Valpara\u00edso",
    "Conc\u00f3n": "Valpara\u00edso",
    "Vina del Mar": "Valpara\u00edso",
    "Vi\u00f1a del Mar": "Valpara\u00edso",
    "Valparaiso": "Valpara\u00edso",
    "Valpara\u00edso": "Valpara\u00edso",
    "Santiago": "Regi\u00f3n Metropolitana",
    "Puente Alto": "Regi\u00f3n Metropolitana",
    "San Jose de Maipo": "Regi\u00f3n Metropolitana",
    "San Jos\u00e9 de Maipo": "Regi\u00f3n Metropolitana",
    "Litueche": "O'Higgins",
    "Concepcion": "Biob\u00edo",
    "Concepci\u00f3n": "Biob\u00edo",
    "Temuco": "Araucan\u00eda",
    "Puerto Montt": "Los Lagos",
    "Coyhaique": "Ays\u00e9n",
    "Punta Arenas": "Magallanes",
}


def registrar_aviso(driver, peor):
    """Llena el formulario de avisos.php con la estacion mas critica."""

    print("\n[4] Abriendo el modulo de avisos mediante un clic...")
    WebDriverWait(driver, TIEMPO_MAXIMO).until(
        EC.element_to_be_clickable((By.CSS_SELECTOR, 'a[href="avisos.php"]'))
    ).click()
    WebDriverWait(driver, TIEMPO_MAXIMO).until(EC.url_contains("/avisos.php"))

    # Espera a que el formulario este disponible antes de tocarlo
    WebDriverWait(driver, TIEMPO_MAXIMO).until(
        EC.element_to_be_clickable((By.NAME, "region"))
    )
    esperar(PAUSA_PASO, "que el formulario quede listo")

    # CONDICIONAL 3: si la ciudad no esta en el diccionario se usa un valor
    # por defecto, en vez de dejar caer el robot.
    if peor["estacion"] in REGIONES:
        region = REGIONES[peor["estacion"]]
    else:
        region = "Regi\u00f3n Metropolitana"
        print("   Aviso: '" + peor["estacion"] + "' sin region mapeada, se usa " + region)

    # CONDICIONAL 4: el tipo de alerta depende de que tan alto este el MP 2.5
    if peor["calidad_aire"] == "CRITICO":
        tipo_alerta = "Alerta Roja"
    elif peor["calidad_aire"] == "MODERADO":
        tipo_alerta = "Alerta Amarilla"
    else:
        tipo_alerta = "Alerta Temprana Preventiva"

    descripcion = (
        "[RPA MeteoStream] Estacion " + peor["estacion"] + " (" + peor["zona"] + "): "
        + "MP 2.5 = " + str(peor["mp25"]) + " ug/m3 -> calidad " + peor["calidad_aire"] + ". "
        + "Temperatura " + str(peor["temperatura_c"]) + " C, "
        + "humedad " + str(peor["humedad_pct"]) + "%, "
        + "condicion " + peor["condicion"] + ". "
        + "Registro automatico " + datetime.now().strftime("%d-%m-%Y %H:%M") + "."
    )

    # --- Manejo de elementos web: dos <select> y un <textarea> ---
    # Se deja una pausa entre campo y campo para que se vea el llenado.
    Select(driver.find_element(By.NAME, "region")).select_by_visible_text(region)
    print("   Region      : " + region)
    time.sleep(1)

    Select(driver.find_element(By.NAME, "tipo_alerta")).select_by_visible_text(tipo_alerta)
    print("   Tipo alerta : " + tipo_alerta)
    time.sleep(1)

    caja_texto = driver.find_element(By.NAME, "descripcion")
    caja_texto.clear()
    caja_texto.send_keys(descripcion)
    print("   Descripcion : " + descripcion[:70] + "...")

    esperar(PAUSA_PASO, "revisar el formulario antes de enviarlo")
    capturar(driver, "03_formulario_lleno")

    # CONDICIONAL 5: se envia o no segun la configuracion de arriba
    if PUBLICAR_AVISO:
        driver.find_element(By.CSS_SELECTOR, "form button[type='submit']").click()
        WebDriverWait(driver, TIEMPO_MAXIMO).until(
            EC.presence_of_element_located((By.NAME, "region"))
        )
        esperar(PAUSA_PASO, "que el sitio confirme el registro")
        print("   Aviso publicado.")
        capturar(driver, "04_aviso_publicado")
    else:
        print("   Modo simulacion: el formulario quedo lleno pero NO se envio.")


# ---------------------------------------------------------------
# 5. PROGRAMA PRINCIPAL
# ---------------------------------------------------------------

def main():
    driver = None
    inicio = time.time()
    try:
        print("=" * 60)
        print(" RPA METEOSTREAM - navegador: " + NAVEGADOR)
        print(" Evidencias en: " + CARPETA)
        print("=" * 60)

        # Validar el maestro antes de abrir el navegador.
        umbrales = cargar_umbrales(ARCHIVO_UMBRALES)
        driver = crear_driver(NAVEGADOR)

        # Acciones 1 y 2
        estaciones = extraer_estaciones(driver)
        descargar_telemetria(driver)
        consultar_sismos(driver)

        # Validacion: si el sitio no entrego datos, no tiene sentido seguir
        if len(estaciones) == 0:
            print("\nNo se extrajo ninguna estacion. Proceso detenido.")
            return

        for estacion in estaciones:
            estacion["estado_alerta"] = evaluar_alerta(estacion, umbrales)
        guardar_csv(estaciones)
        guardar_excel(estaciones, umbrales, CARPETA, URL_PORTADA)

        # Se busca la estacion con peor calidad de aire.
        # Las que no tienen dato (None) se descartan con un filtro.
        con_dato = [e for e in estaciones if e["mp25"] is not None]

        if len(con_dato) == 0:
            print("\nNinguna estacion informo MP 2.5, no se registra aviso.")
            return

        peor = max(con_dato, key=lambda e: e["mp25"])
        print("\n   Estacion mas critica: " + peor["estacion"]
              + " (" + str(peor["mp25"]) + " ug/m3 - " + peor["calidad_aire"] + ")")

        # Accion 4
        registrar_aviso(driver, peor)

        print("\n" + "=" * 60)
        print(" PROCESO COMPLETADO SIN ERRORES")
        print(" Duracion: " + str(round(time.time() - inicio, 1)) + " segundos")
        print("=" * 60)

    # --- Manejo de excepciones: cada error tiene su mensaje claro ---
    except TimeoutException:
        print("\nERROR: el sitio demoro mas de " + str(TIEMPO_MAXIMO) + "s en cargar.")
        print("Suba el valor de TIEMPO_MAXIMO en la configuracion y reintente.")
        if driver:
            capturar(driver, "error_timeout")

    except NoSuchElementException as e:
        print("\nERROR: no se encontro un elemento esperado en la pagina.\n" + str(e))
        if driver:
            capturar(driver, "error_elemento")

    except WebDriverException as e:
        print("\nERROR del navegador o del driver.\n" + str(e))

    except Exception as e:
        print("\nERROR inesperado: " + type(e).__name__ + " -> " + str(e))

    finally:
        # El finally se ejecuta siempre: haya error o no, se cierra el navegador
        if driver:
            esperar(PAUSA_PASO, "dejar ver el resultado final antes de cerrar")
            descargas = getattr(driver, "telemetria_descargas", None)
            driver.quit()
            if descargas and os.path.exists(descargas):
                import shutil
                shutil.rmtree(descargas, ignore_errors=True)
            print("\nNavegador cerrado.")
            
        print("\n" + "=" * 60)
        print(" [MARCA DE AGUA] - AUTOMATIZACION RPA")
        print(" Hecho por Fabiola villagra Ignacio gonzales y Francisca Sanhueza")
        print("=" * 60 + "\n")


if __name__ == "__main__":
    main()
