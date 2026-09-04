"""Generación de reportes locales; no realiza ninguna interacción web."""

from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from .models import Earthquake, StationMeasurement


HEADER_FILL = PatternFill("solid", fgColor="17365D")
HEADER_FONT = Font(color="FFFFFF", bold=True)
RED_FILL = PatternFill("solid", fgColor="F4CCCC")
YELLOW_FILL = PatternFill("solid", fgColor="FFF2CC")


def _style_sheet(sheet) -> None:
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for cell in sheet[1]:
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
    for column_cells in sheet.columns:
        longest = max(len(str(cell.value or "")) for cell in column_cells)
        sheet.column_dimensions[get_column_letter(column_cells[0].column)].width = min(longest + 3, 55)


def export_excel(
    stations: list[StationMeasurement],
    earthquakes: list[Earthquake],
    notices: list[str],
    destination: Path,
) -> Path:
    workbook = Workbook()
    climate = workbook.active
    climate.title = "Clima"
    climate.append(
        ["Zona", "Ciudad", "Temp °C", "Humedad %", "PM2.5", "Condición", "Máx °C", "Mín °C", "Lluvia mm", "Alerta", "Motivo"]
    )
    for item in stations:
        climate.append(
            [item.zone, item.city, item.temperature, item.humidity, item.pm25, item.condition,
             item.maximum, item.minimum, item.rainfall, item.alert_level, item.alert_reason]
        )
    if climate.max_row >= 2:
        climate.conditional_formatting.add(
            f"A2:K{climate.max_row}", FormulaRule(formula=["$J2=\"ROJO\""], fill=RED_FILL)
        )
        climate.conditional_formatting.add(
            f"A2:K{climate.max_row}", FormulaRule(formula=["$J2=\"AMARILLO\""], fill=YELLOW_FILL)
        )
    _style_sheet(climate)

    seismic = workbook.create_sheet("Sismos")
    seismic.append(["Alerta origen", "Magnitud", "Ubicación", "Fecha y hora", "Profundidad", "Alerta por umbral"])
    for item in earthquakes:
        seismic.append([item.source_alert, item.magnitude, item.location, item.occurred_at, item.depth, "SÍ" if item.threshold_alert else "NO"])
    _style_sheet(seismic)

    notice_sheet = workbook.create_sheet("Avisos")
    notice_sheet.append(["Avisos activos observados"])
    for notice in notices:
        notice_sheet.append([notice])
    _style_sheet(notice_sheet)

    destination.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(destination)
    return destination


def export_json(
    stations: list[StationMeasurement],
    earthquakes: list[Earthquake],
    notices: list[str],
    forecast_summary: str,
    destination: Path,
) -> Path:
    payload = {
        "timestamp": datetime.now().astimezone().isoformat(timespec="seconds"),
        "status": "SUCCESS",
        "estaciones_procesadas": len(stations),
        "alertas_meteorologicas": sum(s.alert_level != "VERDE" for s in stations),
        "sismos_procesados": len(earthquakes),
        "pronostico": forecast_summary,
        "estaciones": [item.to_dict() for item in stations],
        "sismos": [item.to_dict() for item in earthquakes],
        "avisos": notices,
    }
    destination.parent.mkdir(parents=True, exist_ok=True)
    destination.write_text(json.dumps(payload, ensure_ascii=False, indent=2), encoding="utf-8")
    return destination

